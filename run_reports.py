"""
Step 2 - Run the reports.

Runs the twelve SQL queries in sql/reports.sql against data/sales.db and
writes each result set into a formatted Excel workbook.

Run:  python run_reports.py
"""

import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
DB_PATH = BASE / "data" / "sales.db"
SQL_PATH = BASE / "sql" / "reports.sql"
OUT_PATH = BASE / "reports" / "sales_report.xlsx"

HEADER_FILL = PatternFill("solid", start_color="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def read_queries(path):
    """Split reports.sql into (name, sql) pairs.

    Each query begins with a line like:  -- name: revenue_by_region
    """
    queries, name, buffer = [], None, []

    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip().lower().startswith("-- name:"):
            if name and buffer:
                queries.append((name, "\n".join(buffer).strip()))
            name = line.split(":", 1)[1].strip()
            buffer = []
        elif not line.strip().startswith("--") or buffer:
            buffer.append(line)

    if name and buffer:
        queries.append((name, "\n".join(buffer).strip()))

    return queries


def tidy_sheet(worksheet):
    """Sensible column widths and a frozen, styled header row."""
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.freeze_panes = "A2"

    for column_cells in worksheet.columns:
        longest = max(len(str(cell.value or "")) for cell in column_cells)
        letter = get_column_letter(column_cells[0].column)
        worksheet.column_dimensions[letter].width = min(max(longest + 3, 11), 34)


def main():
    if not DB_PATH.exists():
        print("data/sales.db not found. Run:  python setup_database.py")
        return

    OUT_PATH.parent.mkdir(exist_ok=True)
    connection = sqlite3.connect(DB_PATH)
    queries = read_queries(SQL_PATH)
    print(f"Found {len(queries)} queries in {SQL_PATH.name}\n")

    results = []
    for name, sql in queries:
        try:
            frame = pd.read_sql_query(sql, connection)
            results.append((name, frame))
            print(f"OK      {name:<32} {len(frame):>4} rows")
        except Exception as error:
            print(f"FAILED  {name:<32} {error}")
    connection.close()

    if not results:
        print("\nNo queries succeeded, so no report was written.")
        return

    try:
        with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
            for name, frame in results:
                sheet = name[:31]
                frame.to_excel(writer, sheet_name=sheet, index=False)
                tidy_sheet(writer.sheets[sheet])
    except PermissionError:
        print(f"\nCould not write {OUT_PATH.name} - close it in Excel and run again.")
        return

    print(f"\nReport written to {OUT_PATH}")


if __name__ == "__main__":
    main()
